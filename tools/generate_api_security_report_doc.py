from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Auditoria_seguridad_APIs_app_movil_CACEI.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.allow_autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_title(doc):
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Investigacion documental y guia de auditoria de seguridad")
    run.font.name = "Calibri"
    run.font.size = Pt(20)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Herramientas DAST y SAST para APIs, backend y aplicaciones moviles")
    run.font.name = "Calibri"
    run.font.size = Pt(12)
    run.italic = True
    run.font.color.rgb = RGBColor.from_string("1F4D78")

    meta = doc.add_table(rows=4, cols=2)
    set_table_widths(meta, [1.7, 4.8])
    rows = [
        ("Proyecto integrador", "CACEI Tutorias - App movil + API Backend"),
        ("Equipo", "____________________________________________"),
        ("Fecha", "____________________________________________"),
        ("Asignatura / docente", "____________________________________________"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        set_cell_text(row.cells[0], label, bold=True)
        set_cell_text(row.cells[1], value)
        set_cell_shading(row.cells[0], "F2F4F7")


def add_paragraph(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_source_note(doc, text):
    p = doc.add_paragraph(text)
    p.style = doc.styles["Normal"]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    for run in p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("555555")


def add_comparison_table(doc):
    doc.add_heading("Cuadro comparativo de herramientas", level=1)
    add_source_note(
        doc,
        "Sintesis elaborada con base en documentacion oficial de ZAP, ProjectDiscovery Nuclei, MobSF, Bearer CLI y OWASP API Security Top 10.",
    )
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = [
        "Herramienta",
        "Tipo",
        "Objeto de analisis",
        "Fortalezas",
        "Limitaciones",
        "Instalacion sugerida",
    ]
    widths = [1.0, 0.8, 1.15, 1.45, 1.35, 0.75]
    set_table_widths(table, widths)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color="0B2545")
        set_cell_shading(cell, "E8EEF5")

    rows = [
        [
            "OWASP ZAP",
            "DAST",
            "APIs activas definidas por OpenAPI, SOAP o GraphQL.",
            "Importa definiciones OpenAPI, ejecuta escaneo activo y pasivo, genera reportes HTML/JSON/XML y permite configurar reglas por severidad.",
            "No reemplaza pruebas manuales de autorizacion por rol; puede generar ruido si la API no tiene datos de prueba controlados.",
            "Docker",
        ],
        [
            "Nuclei",
            "DAST basado en plantillas",
            "Endpoints, hosts, OpenAPI/Swagger, servicios e infraestructura expuesta.",
            "Templates YAML comunitarios, filtros por tags/severidad, CI/CD, alta velocidad y pruebas especificas reproducibles.",
            "La calidad depende de las plantillas elegidas; requiere definir alcance para evitar pruebas invasivas.",
            "Docker / binario",
        ],
        [
            "MobSF",
            "SAST movil + DAST movil",
            "APK, IPA o app movil instrumentada.",
            "Ingenieria inversa estatica, deteccion de URLs, llaves embebidas, permisos, configuracion de red y debilidades de plataforma movil.",
            "El analisis dinamico requiere entorno/emulador; algunos hallazgos necesitan validacion manual.",
            "Docker",
        ],
        [
            "Bearer",
            "SAST backend",
            "Repositorio de API/backend.",
            "Detecta patrones inseguros, secretos, exposicion de datos sensibles, problemas de privacidad y genera reportes de seguridad, privacidad o SARIF.",
            "Su cobertura depende del lenguaje y framework; no observa comportamiento real en produccion.",
            "CLI / Docker",
        ],
    ]
    for data in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, data):
            set_cell_text(cell, value)


def add_theory(doc):
    doc.add_heading("1. Marco teorico", level=1)
    add_paragraph(
        doc,
        "La seguridad de APIs en aplicaciones moviles debe evaluarse desde dos frentes complementarios: el comportamiento observable de los servicios desplegados y la calidad del codigo fuente o binario que consume dichos servicios. El analisis dinamico de seguridad (DAST) prueba una API en ejecucion, enviando solicitudes reales para detectar fallos de configuracion, inyeccion, exposicion de informacion o respuestas inseguras. El analisis estatico de seguridad (SAST) revisa codigo fuente o binarios sin ejecutar el sistema, buscando patrones inseguros antes del despliegue.",
    )
    add_paragraph(
        doc,
        "En el contexto movil, el riesgo no se limita al backend. Una aplicacion puede exponer URLs internas, tokens, llaves, configuraciones debiles de transporte o permisos excesivos. Por ello, una auditoria robusta combina DAST sobre endpoints activos, SAST sobre el backend y analisis estatico del instalador movil.",
    )

    doc.add_heading("1.1 OWASP API Security Top 10 como referencia", level=2)
    add_paragraph(
        doc,
        "OWASP API Security Top 10 2023 organiza los riesgos mas relevantes para APIs modernas. Entre ellos se encuentran autorizacion rota a nivel de objeto, autenticacion rota, autorizacion incorrecta a nivel de propiedades, consumo irrestricto de recursos, autorizacion rota a nivel de funcion, flujos de negocio sensibles expuestos, SSRF, mala configuracion de seguridad, inventario deficiente y consumo inseguro de APIs de terceros.",
    )
    add_source_note(
        doc,
        "Fuente: OWASP API Security Top 10 2023, https://owasp.org/API-Security/editions/2023/en/0x11-t10/",
    )

    doc.add_heading("1.2 SAST, DAST y auditoria movil", level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_widths(table, [1.0, 1.6, 2.1, 1.8])
    for cell, text in zip(
        table.rows[0].cells,
        ["Categoria", "Momento de uso", "Que detecta mejor", "Herramientas del entregable"],
    ):
        set_cell_text(cell, text, bold=True, color="0B2545")
        set_cell_shading(cell, "E8EEF5")
    rows = [
        [
            "SAST",
            "Antes del despliegue o en CI/CD.",
            "Secretos, patrones inseguros, mala gestion de datos, fallos repetibles en codigo.",
            "Bearer, MobSF estatico.",
        ],
        [
            "DAST",
            "Con API desplegada en staging o produccion controlada.",
            "Errores de configuracion, respuestas inseguras, inyeccion, cabeceras, endpoints expuestos.",
            "OWASP ZAP, Nuclei.",
        ],
        [
            "Analisis movil",
            "Sobre APK/IPA y, si aplica, entorno dinamico.",
            "URLs embebidas, llaves hardcoded, permisos, configuracion de red, debilidades de plataforma.",
            "MobSF.",
        ],
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)


def add_tools(doc):
    doc.add_heading("2. Herramientas investigadas", level=1)

    doc.add_heading("2.1 OWASP ZAP", level=2)
    add_paragraph(
        doc,
        "OWASP ZAP es una herramienta DAST que puede funcionar como proxy de interceptacion y como escaner automatizado. Para APIs, su script zap-api-scan.py esta orientado a definiciones OpenAPI, SOAP o GraphQL: importa la especificacion, identifica las URLs y ejecuta reglas activas y pasivas ajustadas para servicios API.",
    )
    add_bullets(
        doc,
        [
            "Uso principal: escaneo automatizado de APIs publicadas en staging o produccion controlada.",
            "Entrada habitual: archivo o URL OpenAPI/Swagger.",
            "Salida: reportes HTML, Markdown, XML o JSON.",
            "Riesgos detectables: inyeccion, mala configuracion de cabeceras, informacion sensible en URL o respuestas, metodos inseguros, errores 5xx, path traversal, XXE y otros patrones cubiertos por sus reglas.",
        ],
    )
    add_source_note(doc, "Fuente: ZAP API Scan, https://www.zaproxy.org/docs/docker/api-scan/")

    doc.add_heading("2.2 Nuclei", level=2)
    add_paragraph(
        doc,
        "Nuclei es un escaner rapido basado en templates YAML. Cada plantilla define solicitudes, condiciones de coincidencia, severidad y metadatos para detectar una debilidad concreta. Esta arquitectura permite automatizar pruebas repetibles, crear templates propios y ejecutar paquetes comunitarios contra endpoints activos.",
    )
    add_bullets(
        doc,
        [
            "Uso principal: pruebas DAST especificas, validacion de CVEs, exposiciones, configuraciones inseguras y endpoints conocidos.",
            "Entrada: URL individual, listas de URLs, hosts, OpenAPI o Swagger mediante input-mode.",
            "Ventaja clave: filtros por tags, severidad, workflows y ejecucion en CI/CD.",
            "Riesgo operativo: debe limitarse a dominios autorizados y ambientes controlados.",
        ],
    )
    add_source_note(
        doc,
        "Fuentes: ProjectDiscovery Nuclei Overview, Install y Running, https://docs.projectdiscovery.io/opensource/nuclei/overview",
    )

    doc.add_heading("2.3 MobSF", level=2)
    add_paragraph(
        doc,
        "MobSF es un framework de seguridad movil para Android, iOS y Windows que automatiza analisis estatico y dinamico. En analisis estatico procesa APK, IPA u otros paquetes moviles para identificar permisos peligrosos, endpoints embebidos, llaves codificadas, configuracion insegura de red, uso de criptografia debil y otros indicadores de riesgo.",
    )
    add_bullets(
        doc,
        [
            "Uso principal: auditoria del instalador movil antes de distribuirlo.",
            "Entradas: APK o IPA generados por el equipo.",
            "Salidas: reporte con severidad, evidencias, permisos, URLs, trackers, secretos y configuraciones relevantes.",
            "Complemento recomendado: validar manualmente los hallazgos para separar falsos positivos de riesgos explotables.",
        ],
    )
    add_source_note(
        doc,
        "Fuente: repositorio oficial MobSF, https://github.com/MobSF/Mobile-Security-Framework-MobSF",
    )

    doc.add_heading("2.4 Bearer", level=2)
    add_paragraph(
        doc,
        "Bearer CLI es una herramienta SAST enfocada en riesgos de seguridad y privacidad dentro del codigo fuente. Su comando principal bearer scan revisa un proyecto local y reporta hallazgos por severidad, ubicacion de archivo y regla aplicada. Tambien permite reportes de privacidad, dataflow, SARIF, JSON, YAML y HTML.",
    )
    add_bullets(
        doc,
        [
            "Uso principal: revision del repositorio backend o API antes de desplegar.",
            "Detecta: patrones de inyeccion, secretos, almacenamiento inseguro, filtrado de datos sensibles y malas practicas de privacidad segun lenguaje soportado.",
            "Configuracion: permite limitar severidades, omitir reglas justificadas, generar reportes y usar diferencial por rama.",
            "Limitacion: no valida el comportamiento real de la API desplegada; se complementa con DAST.",
        ],
    )
    add_source_note(doc, "Fuente: Bearer CLI Quick Start, https://docs.bearer.com/quickstart/")


def add_installation(doc):
    doc.add_heading("3. Flujo de instalacion sugerido con Docker o CLI", level=1)
    add_paragraph(
        doc,
        "Los siguientes comandos son plantillas de ejecucion. Deben ajustarse con las rutas reales del proyecto, el APK final, la URL de staging/produccion y, si aplica, tokens de autenticacion de prueba.",
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_widths(table, [1.0, 2.75, 2.75])
    for cell, text in zip(table.rows[0].cells, ["Herramienta", "Instalacion", "Ejecucion sugerida"]):
        set_cell_text(cell, text, bold=True, color="0B2545")
        set_cell_shading(cell, "E8EEF5")
    rows = [
        [
            "OWASP ZAP",
            "Usar imagen Docker oficial de ZAP.",
            "docker run -t ghcr.io/zaproxy/zaproxy:stable zap-api-scan.py -t https://api.ejemplo.com/openapi.json -f openapi -r zap-report.html -J zap-report.json",
        ],
        [
            "Nuclei",
            "docker pull projectdiscovery/nuclei:latest",
            "docker run --rm projectdiscovery/nuclei:latest -u https://api.ejemplo.com -severity critical,high,medium",
        ],
        [
            "MobSF",
            "Levantar contenedor MobSF y abrir la interfaz web.",
            "docker run -it --rm -p 8000:8000 opensecurity/mobile-security-framework-mobsf:latest",
        ],
        [
            "Bearer",
            "CLI: curl -sfL https://raw.githubusercontent.com/Bearer/bearer/main/contrib/install.sh | sh",
            "bearer scan ./backend --format html --output bearer-report.html",
        ],
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)


def add_practical_phase(doc):
    doc.add_heading("4. Fase 2: Aplicacion practica en el proyecto CACEI", level=1)
    add_paragraph(
        doc,
        "Esta seccion esta disenada para documentar la auditoria aplicada directamente al proyecto integrador CACEI. El alcance practico contempla la aplicacion movil CACEI, el repositorio de la API/backend CACEI y los endpoints desplegados en el ambiente autorizado de staging o produccion controlada. El objetivo no es solo mostrar capturas, sino justificar tecnicamente que se comprendieron los hallazgos y se definieron acciones de mitigacion.",
    )

    doc.add_heading("4.1 Bitacora de ejecucion", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    set_table_widths(table, [1.0, 1.25, 1.4, 1.45, 1.4])
    headers = ["Herramienta", "Objetivo auditado", "Comando / version", "Evidencia", "Resultado general"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color="0B2545")
        set_cell_shading(cell, "E8EEF5")
    rows = [
        [
            "MobSF",
            "Instalador APK/IPA de CACEI Tutorías",
            "Version de MobSF y hash del instalador",
            "Captura del reporte MobSF y resumen PDF/HTML",
            "Pendiente",
        ],
        [
            "Bearer",
            "Repositorio de la API/backend CACEI",
            "bearer scan ./backend",
            "Captura del reporte y archivo de salida",
            "Pendiente",
        ],
        [
            "OWASP ZAP",
            "OpenAPI/Swagger de la API CACEI desplegada",
            "zap-api-scan.py contra la especificacion OpenAPI",
            "Captura del escaneo y reporte HTML/JSON",
            "Pendiente",
        ],
        [
            "Nuclei",
            "Endpoints activos autorizados de la API CACEI",
            "nuclei -u <api-cacei> -severity critical,high,medium",
            "Captura de consola y reporte exportado",
            "Pendiente",
        ],
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            set_cell_text(cell, value)

    doc.add_heading("4.2 Inventario de hallazgos", level=2)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    set_table_widths(table, [0.65, 1.05, 1.0, 1.35, 1.65, 0.8])
    headers = ["ID", "Herramienta", "Criticidad", "Descripcion", "Evidencia tecnica", "Estado"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color="0B2545")
        set_cell_shading(cell, "E8EEF5")
    for index in range(1, 6):
        cells = table.add_row().cells
        values = [
            f"H-{index:02d}",
            "Pendiente",
            "Alta/Media/Baja",
            "Debilidad detectada en app movil, backend o API CACEI",
            "Archivo, endpoint, captura o regla asociada",
            "Abierto",
        ]
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)

    doc.add_heading("4.3 Plan de mitigacion", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    set_table_widths(table, [0.75, 1.7, 1.6, 1.2, 1.25])
    headers = ["Hallazgo", "Correccion propuesta", "Cambio en codigo/configuracion", "Responsable", "Fecha objetivo"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_text(cell, text, bold=True, color="0B2545")
        set_cell_shading(cell, "E8EEF5")
    for index in range(1, 6):
        cells = table.add_row().cells
        values = [f"H-{index:02d}", "Pendiente", "Pendiente", "Pendiente", "Pendiente"]
        for cell, value in zip(cells, values):
            set_cell_text(cell, value)

    doc.add_heading("4.4 Justificacion si no se detectan vulnerabilidades", level=2)
    add_paragraph(
        doc,
        "Si alguna herramienta no arroja vulnerabilidades en CACEI, el equipo debe documentar que buenas practicas redujeron la exposicion: autenticacion robusta, autorizacion por rol, validacion de entrada, no inclusion de secretos en el cliente movil, uso de HTTPS, gestion segura de tokens, configuracion segura de CORS, rate limiting, manejo de errores sin datos sensibles, inventario actualizado de endpoints CACEI y separacion de ambientes.",
    )


def add_execution_plan(doc):
    doc.add_heading("5. Procedimiento recomendado de auditoria", level=1)
    add_numbered(
        doc,
        [
            "Definir alcance: URL base de API, endpoints permitidos, ambiente de staging o produccion controlada, credenciales de prueba y ventana de ejecucion.",
            "Generar artefactos: APK/IPA final de la app movil, repositorio backend actualizado y especificacion OpenAPI/Swagger.",
            "Ejecutar MobSF contra el instalador movil y exportar reporte.",
            "Ejecutar Bearer sobre el repositorio backend y guardar reporte HTML/SARIF/JSON.",
            "Ejecutar OWASP ZAP API Scan usando la especificacion OpenAPI y guardar reportes HTML y JSON.",
            "Ejecutar Nuclei sobre la API con templates relevantes, priorizando severidades critical, high y medium.",
            "Clasificar hallazgos por criticidad y eliminar falsos positivos con justificacion tecnica.",
            "Definir mitigaciones, responsables y fecha de correccion antes del despliegue final.",
        ],
    )

    doc.add_heading("5.1 Buenas practicas de seguridad durante las pruebas", level=2)
    add_bullets(
        doc,
        [
            "No ejecutar escaneos agresivos contra sistemas de terceros o dominios no autorizados.",
            "Usar cuentas de prueba sin datos personales reales.",
            "Guardar reportes con control de acceso, ya que pueden contener endpoints, rutas internas o fragmentos sensibles.",
            "Programar pruebas DAST en horarios controlados para evitar saturar la API.",
            "Registrar versiones de herramientas, fecha, comandos y ambiente para que la auditoria sea reproducible.",
        ],
    )


def add_sources(doc):
    doc.add_heading("6. Fuentes consultadas", level=1)
    sources = [
        "OWASP API Security Top 10 2023: https://owasp.org/API-Security/editions/2023/en/0x11-t10/",
        "ZAP API Scan: https://www.zaproxy.org/docs/docker/api-scan/",
        "ProjectDiscovery Nuclei Overview: https://docs.projectdiscovery.io/opensource/nuclei/overview",
        "ProjectDiscovery Nuclei Install: https://docs.projectdiscovery.io/opensource/nuclei/install",
        "ProjectDiscovery Nuclei Running: https://docs.projectdiscovery.io/opensource/nuclei/running",
        "MobSF official repository: https://github.com/MobSF/Mobile-Security-Framework-MobSF",
        "Bearer CLI Quick Start: https://docs.bearer.com/quickstart/",
        "Bearer CLI Configure Scan: https://docs.bearer.com/guides/configure-scan/",
    ]
    add_bullets(doc, sources)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("CACEI Tutorias | Auditoria de seguridad API y movil")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("555555")


def build():
    doc = Document()
    style_document(doc)
    add_title(doc)
    add_paragraph(
        doc,
        "Este documento integra la investigacion documental de herramientas abiertas para auditoria de seguridad y una guia practica para aplicarlas al proyecto integrador. La estructura cubre herramientas DAST, SAST y analisis movil, con enfasis en APIs, backend y aplicacion movil.",
    )
    add_theory(doc)
    add_tools(doc)
    add_comparison_table(doc)
    add_installation(doc)
    add_practical_phase(doc)
    add_execution_plan(doc)
    add_sources(doc)
    add_footer(doc)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
